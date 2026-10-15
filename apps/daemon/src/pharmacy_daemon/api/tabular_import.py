"""Tabular data parser and sample template generator for CSV and Excel (.xlsx) bulk imports."""

import csv
from datetime import datetime, timezone
import io
import re
from typing import Any, Dict, List, Optional
import uuid

import openpyxl


# In-memory session cache for previewing import batches before execution
# Keys: staging_token (str) -> dict(created_at: datetime, type: str, rows: List[dict])
_STAGED_IMPORT_SESSIONS: Dict[str, Dict[str, Any]] = {}


def save_staged_import(import_type: str, rows: List[Dict[str, Any]]) -> str:
    """Store validated preview rows in an ephemeral staging session and return a token."""
    # Evict sessions older than 30 minutes
    now = datetime.now(timezone.utc)
    expired = [
        k for k, v in _STAGED_IMPORT_SESSIONS.items()
        if (now - v["created_at"]).total_seconds() > 1800
    ]
    for k in expired:
        _STAGED_IMPORT_SESSIONS.pop(k, None)

    token = f"stg-{uuid.uuid4().hex[:16]}"
    _STAGED_IMPORT_SESSIONS[token] = {
        "created_at": now,
        "type": import_type,
        "rows": rows,
    }
    return token


def get_staged_import(token: str) -> Optional[Dict[str, Any]]:
    """Retrieve an active staged import session by token."""
    session = _STAGED_IMPORT_SESSIONS.get(token)
    if not session:
        return None
    # Check expiry
    if (datetime.now(timezone.utc) - session["created_at"]).total_seconds() > 1800:
        _STAGED_IMPORT_SESSIONS.pop(token, None)
        return None
    return session


def clear_staged_import(token: str) -> None:
    """Remove a staged import session upon execution or cancellation."""
    _STAGED_IMPORT_SESSIONS.pop(token, None)


def normalize_key(key: str) -> str:
    """Normalize a column header string: lowercase, remove non-alphanumeric symbols."""
    return re.sub(r"[^a-z0-9]", "", str(key).strip().lower())


def find_col(row: Dict[str, Any], candidates: List[str], default: Any = None) -> Any:
    """Fuzzy-match and extract a column value from a row using candidate header names."""
    normalized_row = {normalize_key(k): v for k, v in row.items()}
    for cand in candidates:
        norm_cand = normalize_key(cand)
        if norm_cand in normalized_row:
            val = normalized_row[norm_cand]
            if val is not None and str(val).strip() != "":
                return val
    return default


def parse_tabular_upload(file_bytes: bytes, filename: str) -> List[Dict[str, Any]]:
    """Parse raw bytes of a CSV or Excel (.xlsx) file into a list of row dictionaries."""
    lower_name = filename.lower()
    if lower_name.endswith((".xlsx", ".xlsm")):
        return _parse_excel(file_bytes)
    # Default to CSV parser
    return _parse_csv(file_bytes)


def _parse_excel(file_bytes: bytes) -> List[Dict[str, Any]]:
    """Extract rows from the first sheet of an Excel workbook using openpyxl."""
    wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []

    # First non-empty row is treated as header
    header_idx = 0
    for idx, r in enumerate(rows):
        if any(c is not None and str(c).strip() for c in r):
            header_idx = idx
            break

    raw_headers = rows[header_idx]
    headers = [str(c).strip() if c is not None else f"col_{i}" for i, c in enumerate(raw_headers)]

    data_rows = []
    for r in rows[header_idx + 1:]:
        # Skip completely empty rows
        if not any(c is not None and str(c).strip() for c in r):
            continue
        row_dict: Dict[str, Any] = {}
        for h, val in zip(headers, r):
            if h and val is not None:
                row_dict[h] = val
        if row_dict:
            data_rows.append(row_dict)

    return data_rows


def _parse_csv(file_bytes: bytes) -> List[Dict[str, Any]]:
    """Extract rows from a CSV file with automatic UTF-8 BOM and encoding detection."""
    try:
        text = file_bytes.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = file_bytes.decode("latin-1", errors="replace")

    stream = io.StringIO(text)
    # Use Sniffer to check dialect or fallback to standard comma
    try:
        sample = text[:2048]
        dialect = csv.Sniffer().sniff(sample)
    except Exception:
        dialect = csv.excel

    reader = csv.reader(stream, dialect=dialect)
    rows = list(reader)
    if not rows:
        return []

    header_idx = 0
    for idx, r in enumerate(rows):
        if any(c.strip() for c in r if isinstance(c, str)):
            header_idx = idx
            break

    raw_headers = rows[header_idx]
    headers = [str(c).strip() if c else f"col_{i}" for i, c in enumerate(raw_headers)]

    data_rows = []
    for r in rows[header_idx + 1:]:
        if not any(c.strip() for c in r if isinstance(c, str)):
            continue
        row_dict: Dict[str, Any] = {}
        for h, val in zip(headers, r):
            if h:
                row_dict[h] = val.strip() if isinstance(val, str) else val
        if row_dict:
            data_rows.append(row_dict)

    return data_rows


def generate_csv_template(headers: List[str], sample_rows: List[List[Any]]) -> bytes:
    """Generate a clean downloadable CSV sample template with headers and example data."""
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(headers)
    for r in sample_rows:
        writer.writerow(r)
    return output.getvalue().encode("utf-8-sig")


def generate_xlsx_template(sheet_title: str, headers: List[str], sample_rows: List[List[Any]]) -> bytes:
    """Generate a downloadable Excel (.xlsx) template with styled headers."""
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_title[:31]

    # Header style
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    thin_border = Border(
        left=Side(style="thin", color="CCCCCC"),
        right=Side(style="thin", color="CCCCCC"),
        top=Side(style="thin", color="CCCCCC"),
        bottom=Side(style="thin", color="CCCCCC"),
    )

    ws.append(headers)
    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    # Data rows
    for r in sample_rows:
        ws.append(r)

    for row in ws.iter_rows(min_row=2, max_row=len(sample_rows) + 1, min_col=1, max_col=len(headers)):
        for cell in row:
            cell.border = thin_border
            cell.alignment = Alignment(vertical="center")

    # Auto-adjust column widths
    for col in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        col_letter = col[0].column_letter
        ws.column_dimensions[col_letter].width = max(max_len + 4, 14)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
